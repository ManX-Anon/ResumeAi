"""
resume_parser.py

Extracts raw text from uploaded resume files (PDF or DOCX).
Uses pdfplumber as the primary PDF extractor and falls back to PyPDF2
if pdfplumber fails or returns empty text (e.g. certain malformed PDFs).
"""
from __future__ import annotations

import io
import os
from typing import Tuple

import pdfplumber
from PyPDF2 import PdfReader
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class UnsupportedFileTypeError(Exception):
    pass


class EmptyDocumentError(Exception):
    pass


def get_extension(filename: str) -> str:
    return os.path.splitext(filename.lower())[1]


def _extract_pdf_with_pdfplumber(file_bytes: bytes) -> str:
    text_chunks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_chunks.append(page_text)
    return "\n".join(text_chunks).strip()


def _extract_pdf_with_pypdf2(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    text_chunks = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(text_chunks).strip()


def extract_pdf_text(file_bytes: bytes) -> str:
    try:
        text = _extract_pdf_with_pdfplumber(file_bytes)
        if text:
            return text
    except Exception:
        pass

    try:
        text = _extract_pdf_with_pypdf2(file_bytes)
        return text
    except Exception as exc:
        raise EmptyDocumentError(f"Failed to extract text from PDF: {exc}") from exc


def extract_docx_text(file_bytes: bytes) -> str:
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise EmptyDocumentError(f"Failed to open DOCX file: {exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]

    # Also pull text out of tables (common in resumes with skill tables)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts).strip()


def extract_text(filename: str, file_bytes: bytes) -> Tuple[str, str]:
    """
    Returns (extracted_text, file_type) or raises UnsupportedFileTypeError /
    EmptyDocumentError.
    """
    ext = get_extension(filename)

    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. Only PDF and DOCX are supported."
        )

    if ext == ".pdf":
        text = extract_pdf_text(file_bytes)
    else:
        text = extract_docx_text(file_bytes)

    if not text or len(text.strip()) < 20:
        raise EmptyDocumentError(
            "Could not extract meaningful text from this document. "
            "It may be a scanned image or an empty/corrupted file."
        )

    return text, ext.replace(".", "")
